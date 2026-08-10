"""Genera assets/data.js a partir del libro Excel original.

Uso:
    python tools/build_data.py "ruta/Herramienta de diagnostico TRL.xlsm"

Mantiene la herramienta web sincronizada con el instrumento original: si el
Excel cambia (nuevas preguntas, otros textos), se vuelve a ejecutar este script.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

DEFAULT_XLSM = Path(__file__).resolve().parents[2] / "Herramienta de diagnostico TRL.xlsm"
OUT = Path(__file__).resolve().parents[1] / "assets" / "data.js"

# Hoja de preguntas -> (nombre corto, titulo de la seccion, fila inicial)
SECTIONS = [
    ("1. General", "General", "Aspectos básicos del grado de desarrollo de la tecnología", 5),
    ("2. Mercado", "Mercado", "Oportunidad de mercado", 5),
    ("3. Desarrollo", "Desarrollo", "Diseño y desarrollo", 6),
    ("4. Integracion", "Integración", "Integración", 5),
    ("5. Experimentacion", "Experimentación", "Experimentación y validación", 5),
    ("6. Seguridad", "Seguridad", "Seguridad", 5),
    ("7. Produccion", "Producción", "Producción (generación producto/servicio final)", 5),
]

# Umbrales de la hoja RESULTADOS (celdas Z5:AC5) aplicados en la formula de J21.
THRESHOLDS = {9: 0.70, 8: 0.70, 7: 0.70, 6: 0.60, 5: 0.70, 4: 0.60, 3: 0.60, 2: 0.79, 1: 0.79}

MATURITY = {
    0: "Sin nivel alcanzado",
    1: "Concepto",
    2: "Investigación",
    3: "Investigación",
    4: "Prototipo de laboratorio",
    5: "Prototipo de laboratorio",
    6: "Prototipo industrial",
    7: "Prototipo industrial",
    8: "Producción",
    9: "Producción",
}

# Alias adicionales para enlazar el glosario dentro del texto de las preguntas.
ALIASES = {
    "PI": ["PI"],
    "Experimentación a escala ingenieril": ["escala ingenieril", "nivel de ingeniería"],
    "Requisitos de equipamiento": ["pruebas de equipamiento", "requisitos de equipamiento"],
    "Compatibilidad funcional": ["funcionalmente compatible", "compatibilidad funcional"],
    "Documento formal de requisitos del sistema": ["documento formal", "requisitos del sistema"],
    "Experimentación a escala de laboratorio": [
        "escala de laboratorio",
        "nivel de laboratorio",
        "entorno de laboratorio",
    ],
    "Oportunidad de mercado": ["oportunidad de mercado"],
    "Límites operacionales": ["límites operacionales", "límites operaciones"],
    "Nivel operativo máximo": ["nivel operativo previsto máximo", "nivel operativo máximo"],
    "Reivindicaciones de patente": ["reivindicaciones de la patente", "reivindicaciones de patente"],
    "Prueba de concepto": ["prueba de concepto"],
    "Plan de protección": ["plan de protección"],
    "Entorno operativo": ["entorno operacional", "entorno operativo"],
    "Prototipo escalable": ["prototipos escalables", "prototipo escalable"],
    "Entorno de simulación": ["entorno de simulación", "entorno simulado"],
    "Especificaciones del sistema": ["especificaciones del sistema"],
    "Partner": ["partner"],
    "Propuesta de valor": ["propuesta de valor"],
}


# Erratas evidentes del libro original que se corrigen al generar los datos.
# Se listan aquí, y no se editan a mano en data.js, para que sobrevivan a una
# regeneración y quede constancia de qué se cambió respecto de la fuente.
ERRATA = {
    "laboratoriopero todavía": "laboratorio, pero todavía",
    "un plan de protección un plan de protección": "un plan de protección",
    "clientes potenciales identifidados": "clientes potenciales identificados",
}


def clean(value):
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    # El libro original deja espacios sueltos donde había marcas de glosario.
    text = re.sub(r"\s+([,.;:)])", r"\1", text)
    for wrong, right in ERRATA.items():
        text = text.replace(wrong, right)
    return text


def read_questions(wb):
    questions = []
    for sheet, short, title, first_row in SECTIONS:
        ws = wb[sheet]
        items = []
        row = first_row
        while True:
            code = clean(ws.cell(row=row, column=5).value)  # columna E: 1-1), 1-2)...
            text = clean(ws.cell(row=row, column=6).value)  # columna F: enunciado
            trl = ws.cell(row=row, column=2).value          # columna B: TRL asociado
            if not code or not text or trl is None:
                break
            items.append({"id": code.rstrip(")"), "text": text, "trl": int(trl)})
            row += 1
        questions.append({"key": sheet, "short": short, "title": title, "items": items})
    return questions


def read_actions(wb):
    """Plan de trabajo: una accion por pregunta, en el mismo orden del cuestionario."""
    ws = wb["Plan de trabajo"]
    actions = []
    for row in range(4, ws.max_row + 1):
        action = clean(ws.cell(row=row, column=3).value)  # columna C
        if action:
            actions.append(action)
    return actions


def read_levels(wb):
    ws = wb["Guía"]
    levels = []
    stage = ""
    for row in range(13, 24):
        stage = clean(ws.cell(row=row, column=2).value) or stage
        code = clean(ws.cell(row=row, column=3).value)
        if not code:
            continue
        levels.append(
            {
                "trl": int(code.replace("TRL", "").strip()),
                "stage": stage,
                "name": clean(ws.cell(row=row, column=4).value),
                "description": clean(ws.cell(row=row, column=5).value),
            }
        )
    return sorted(levels, key=lambda item: item["trl"])


def read_glossary(wb):
    ws = wb["Glosario"]
    entries = []
    for row in range(6, ws.max_row + 1):
        term = clean(ws.cell(row=row, column=5).value)  # columna E
        definition = clean(ws.cell(row=row, column=6).value)  # columna F
        if term and definition:
            if term == "Límites operaciones":  # errata del libro original
                term = "Límites operacionales"
            entries.append({"term": term, "definition": definition, "aliases": ALIASES.get(term, [])})
    return entries


def read_guide(wb):
    ws = wb["Guía"]
    return [clean(ws.cell(row=row, column=2).value) for row in (5, 6, 7)]


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSM
    wb = openpyxl.load_workbook(src, data_only=True)

    sections = read_questions(wb)
    actions = read_actions(wb)
    total = sum(len(section["items"]) for section in sections)
    if len(actions) != total:
        raise SystemExit(f"Plan de trabajo tiene {len(actions)} acciones y hay {total} preguntas")

    index = 0
    for section in sections:
        for item in section["items"]:
            item["action"] = actions[index]
            index += 1

    data = {
        "sections": sections,
        "levels": read_levels(wb),
        "glossary": read_glossary(wb),
        "guide": read_guide(wb),
        "thresholds": {str(k): v for k, v in THRESHOLDS.items()},
        "maturity": {str(k): v for k, v in MATURITY.items()},
    }

    payload = json.dumps(data, ensure_ascii=False, indent=2)
    OUT.write_text(
        "/* Generado por tools/build_data.py a partir del Excel original. No editar a mano. */\n"
        f"window.TRL_DATA = {payload};\n",
        encoding="utf-8",
    )
    print(f"{OUT} escrito: {total} preguntas, {len(data['glossary'])} términos de glosario")


if __name__ == "__main__":
    main()
